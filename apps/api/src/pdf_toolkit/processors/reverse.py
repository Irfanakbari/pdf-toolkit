import io
import logging
from collections.abc import Callable
from pathlib import Path

import fitz  # PyMuPDF
import openpyxl
import pdfplumber
from pdf2docx import Converter
from pptx import Presentation
from pptx.util import Inches

logger = logging.getLogger(__name__)
ProgressCallback = Callable[[int, int], None]


def pdf_to_word(source: Path, output: Path) -> None:
    cv = Converter(str(source))
    try:
        cv.convert(str(output), start=0, end=None)
    finally:
        cv.close()


def pdf_to_excel(source: Path, output: Path, progress: ProgressCallback | None = None) -> None:
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    has_data = False

    with pdfplumber.open(source) as pdf:
        for idx, page in enumerate(pdf.pages, 1):
            sheet = wb.create_sheet(title=f"Page {idx}")
            tables = page.extract_tables()
            row_cursor = 1

            if tables:
                for table in tables:
                    for row in table:
                        for col_idx, cell in enumerate(row, 1):
                            sheet.cell(row=row_cursor, column=col_idx, value=cell or "")
                        row_cursor += 1
                    row_cursor += 1  # Empty row between tables
                has_data = True
            else:
                # Fallback: extract line-by-line text if no table detected
                text = page.extract_text() or ""
                lines = [line.strip() for line in text.split("\n") if line.strip()]
                for line in lines:
                    sheet.cell(row=row_cursor, column=1, value=line)
                    row_cursor += 1
                if lines:
                    has_data = True
            if progress:
                progress(idx, len(pdf.pages))

    if default_sheet and default_sheet in wb.worksheets:
        wb.remove(default_sheet)

    if not has_data and not wb.worksheets:
        empty_sheet = wb.create_sheet(title="Sheet 1")
        empty_sheet.cell(row=1, column=1, value="No tabular text detected in PDF.")

    wb.save(output)


def pdf_to_powerpoint(source: Path, output: Path, progress: ProgressCallback | None = None) -> None:
    prs = Presentation()
    doc = fitz.open(source)
    try:
        for index, page in enumerate(doc, 1):
            # Render page to image at 150 DPI for balance of sharpness and memory
            pix = page.get_pixmap(dpi=150)
            img_bytes = pix.tobytes("png")

            # Determine slide size matching the PDF aspect ratio
            rect = page.rect
            slide_width = Inches(rect.width / 72.0)
            slide_height = Inches(rect.height / 72.0)

            prs.slide_width = slide_width
            prs.slide_height = slide_height

            blank_layout = prs.slide_layouts[6]  # Blank slide layout
            slide = prs.slides.add_slide(blank_layout)

            image_stream = io.BytesIO(img_bytes)
            slide.shapes.add_picture(
                image_stream, left=0, top=0, width=slide_width, height=slide_height
            )
            if progress:
                progress(index, len(doc))
    finally:
        doc.close()

    prs.save(output)
